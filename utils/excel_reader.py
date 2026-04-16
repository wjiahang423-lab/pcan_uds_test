# ============================================================
# utils/excel_reader.py  ——  UDS 测试用例 Excel 解析层
#
# 对应 Excel 文件：UDS_TestCases_Template.xlsx
# 每个 sheet：第 1 行为合并标题，第 2 行为列名，第 3 行起为数据。
# ============================================================

import os
import re
import openpyxl

_EXCEL_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                           'UDS_TestCases_Template.xlsx')

# ── 内部工具 ────────────────────────────────────────────────

def _parse_hex(val) -> int | None:
    """将 '0x7E0'、'0x01' 等字符串转为 int，空值/'-' 返回 None。"""
    if val is None:
        return None
    s = str(val).strip()
    if s in ('', '-', 'nan', 'NaN', 'N/A'):
        return None
    s = s.split('/')[0].strip()   # 取 '0x12 / 0x31' 中的第一个，仅用于单值场景
    try:
        return int(s, 16)
    except ValueError:
        return None


def _parse_nrc_list(val) -> list[int]:
    """
    将 '0x12 / 0x31' 或 '0x35' 解析为 int 列表。
    空值/'-' 返回空列表。
    """
    if val is None:
        return []
    s = str(val).strip()
    if s in ('', '-', 'nan', 'NaN'):
        return []
    parts = re.split(r'[/,]', s)
    result = []
    for p in parts:
        p = p.strip()
        try:
            result.append(int(p, 16))
        except ValueError:
            pass
    return result


def _parse_bytes(val) -> bytes:
    """将十六进制字符串（如 'AABB CC'）解析为 bytes，空值返回 b''。"""
    if val is None:
        return b''
    s = str(val).strip().replace(' ', '')
    if s in ('', '-', 'nan', 'NaN'):
        return b''
    try:
        return bytes.fromhex(s)
    except ValueError:
        return b''


def _is_case_row(row_id) -> bool:
    """判断一行是否为有效用例行（用例编号形如 XX-NNN）。"""
    if row_id is None:
        return False
    return bool(re.match(r'^[A-Z]+-\d+$', str(row_id).strip()))


def _read_sheet(sheet_name: str) -> list[dict]:
    """
    读取指定 sheet，返回原始 dict 列表（列名为键，保留原始值）。
    跳过空行和非用例行（通过首列判断）。
    """
    wb = openpyxl.load_workbook(_EXCEL_PATH, data_only=True)
    ws = wb[sheet_name]

    # 第 2 行为列名
    headers = [cell.value for cell in ws[2]]

    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if all(v is None for v in row):
            continue
        d = dict(zip(headers, row))
        # 过滤非用例行（如 DTC 末尾的状态位参考子表）
        first_val = row[0]
        if not _is_case_row(first_val):
            continue
        rows.append(d)
    return rows


# ── 公开加载函数 ─────────────────────────────────────────────

def load_diag_session_cases() -> list[dict]:
    """
    读取 DiagSession sheet。
    返回字段：id, name, session_type(int), session_name, pre_session,
              expected(str), nrc(list[int]), note
    """
    raw = _read_sheet('DiagSession')
    result = []
    for r in raw:
        result.append({
            'id':           str(r.get('用例编号', '')).strip(),
            'name':         str(r.get('用例名称', '')).strip(),
            'session_type': _parse_hex(r.get('请求会话类型(hex)')),
            'session_name': str(r.get('会话描述', '')).strip(),
            'pre_session':  str(r.get('前置会话', '')).strip(),
            'expected':     str(r.get('期望响应类型', '')).strip(),
            'nrc':          _parse_nrc_list(r.get('期望 NRC(hex)')),
            'note':         str(r.get('备注', '') or '').strip(),
        })
    return result


def load_ecu_reset_cases() -> list[dict]:
    """
    读取 ECUReset sheet。
    返回字段：id, name, reset_type(int), reset_desc, pre_cond,
              expected(str), nrc(list[int]), note
    """
    raw = _read_sheet('ECUReset')
    result = []
    for r in raw:
        result.append({
            'id':         str(r.get('用例编号', '')).strip(),
            'name':       str(r.get('用例名称', '')).strip(),
            'reset_type': _parse_hex(r.get('复位类型(hex)')),
            'reset_desc': str(r.get('复位类型描述', '')).strip(),
            'pre_cond':   str(r.get('前置条件', '')).strip(),
            'expected':   str(r.get('期望响应类型', '')).strip(),
            'nrc':        _parse_nrc_list(r.get('期望 NRC(hex)')),
            'note':       str(r.get('复位后等待(s) / 备注', '') or '').strip(),
        })
    return result


def load_rdbi_cases() -> list[dict]:
    """
    读取 RDBI sheet。
    返回字段：id, name, did(int), did_desc, session, min_len(int),
              expected_raw(bytes|None), expected(str), nrc(list[int]), note
    """
    raw = _read_sheet('RDBI')
    result = []
    for r in raw:
        raw_hex = r.get('expected_raw (hex，留空=不校验)')
        if raw_hex is None or str(raw_hex).strip() in ('', '-', 'nan', 'NaN'):
            expected_raw = None
        else:
            expected_raw = _parse_bytes(raw_hex)

        min_len_val = r.get('min_len (bytes)')
        try:
            min_len = int(min_len_val) if min_len_val is not None else 0
        except (ValueError, TypeError):
            min_len = 0

        result.append({
            'id':           str(r.get('用例编号', '')).strip(),
            'name':         str(r.get('name（脚本 ID）', '')).strip(),
            'did':          _parse_hex(r.get('DID (hex)')),
            'did_desc':     str(r.get('DID 描述', '')).strip(),
            'session':      str(r.get('所需会话', '')).strip(),
            'min_len':      min_len,
            'expected_raw': expected_raw,
            'expected':     str(r.get('期望响应类型', '')).strip(),
            'nrc':          _parse_nrc_list(r.get('期望 NRC(hex)')),
            'note':         str(r.get('备注', '') or '').strip(),
        })
    return result


def load_security_access_cases() -> list[dict]:
    """
    读取 SecurityAccess sheet。
    返回字段：id, name, seed_level(int), key_level(int|None), pre_session,
              expected(str), nrc(list[int]), algo, note
    """
    raw = _read_sheet('SecurityAccess')
    result = []
    for r in raw:
        result.append({
            'id':          str(r.get('用例编号', '')).strip(),
            'name':        str(r.get('用例名称', '')).strip(),
            'seed_level':  _parse_hex(r.get('SA 级别 Seed(hex)')),
            'key_level':   _parse_hex(r.get('SA 级别 Key(hex)')),
            'pre_session': str(r.get('前置会话', '')).strip(),
            'expected':    str(r.get('期望响应类型', '')).strip(),
            'nrc':         _parse_nrc_list(r.get('期望 NRC(hex)')),
            'algo':        str(r.get('Seed-Key 算法', '') or '').strip(),
            'note':        str(r.get('备注', '') or '').strip(),
        })
    return result


def load_dtc_cases() -> list[dict]:
    """
    读取 DTC sheet，过滤末尾状态位参考子表。
    返回字段：id, name, sid(int), sub_func, pre_cond,
              expected(str), nrc(list[int]), note
    """
    raw = _read_sheet('DTC')
    result = []
    for r in raw:
        result.append({
            'id':       str(r.get('用例编号', '')).strip(),
            'name':     str(r.get('用例名称', '')).strip(),
            'sid':      _parse_hex(r.get('SID')),
            'sub_func': str(r.get('子功能 / 参数', '') or '').strip(),
            'pre_cond': str(r.get('前置条件', '')).strip(),
            'expected': str(r.get('期望响应类型', '')).strip(),
            'nrc':      _parse_nrc_list(r.get('期望 NRC(hex)')),
            'note':     str(r.get('备注', '') or '').strip(),
        })
    return result


def load_routine_cases() -> list[dict]:
    """
    读取 Routine sheet。
    返回字段：id, name, routine_id(int), routine_desc, session,
              start_data(bytes), expect_result_len(int),
              expected(str), nrc(list[int]), result_desc, note
    """
    raw = _read_sheet('Routine')
    result = []
    for r in raw:
        rlen_val = r.get('expect_result_len')
        try:
            expect_result_len = int(rlen_val) if rlen_val is not None else 0
        except (ValueError, TypeError):
            expect_result_len = 0

        result.append({
            'id':                str(r.get('用例编号', '')).strip(),
            'name':              str(r.get('name（脚本 ID）', '')).strip(),
            'routine_id':        _parse_hex(r.get('routine_id (hex)')),
            'routine_desc':      str(r.get('routine 描述', '')).strip(),
            'session':           str(r.get('所需会话', '')).strip(),
            'start_data':        _parse_bytes(r.get('start_data (hex，留空=b\'\')')),
            'expect_result_len': expect_result_len,
            'expected':          str(r.get('期望响应类型', '')).strip(),
            'nrc':               _parse_nrc_list(r.get('期望 NRC(hex)')),
            'result_desc':       str(r.get('结果数据描述', '') or '').strip(),
            'note':              str(r.get('备注', '') or '').strip(),
        })
    return result
